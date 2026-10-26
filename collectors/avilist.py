#!/usr/bin/env python3
"""
Download and convert the AviList checklist.

Downloads the AviList XLSX from the official URL (configured in config.yml)
and converts it to CSV for use by the rest of the pipeline.

Output: the CSV named by avilist.csv_file in config.yml (under raw_data/)

Usage:
    python -m collectors.avilist [--force]

Requires: openpyxl (pip install openpyxl)
"""

import argparse
import csv
import re
import sys
from io import BytesIO
from pathlib import Path
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError

from config import load_config
from collectors._common import ROOT, RAW_DIR, USER_AGENT

# v2026 switched to upper-case, space-separated headers ("SCIENTIFIC NAME").
# Normalise them to the v2025 column names the rest of the pipeline reads.
_HEADER_MAP = {
    "#": "Sequence",
    "RANK": "Taxon_rank",
    "SCIENTIFIC NAME": "Scientific_name",
    "ENGLISH NAME: AVILIST": "English_name_AviList",
    "EXTINCTION STATUS": "Extinct_or_possibly_extinct",
    "IUCN RED LIST CATEGORY": "IUCN_Red_List_Category",
    "BIRDLIFE DATAZONE URL": "BirdLife_DataZone_URL",
    "SPECIES CODE CORNELL LAB": "Species_code_Cornell_Lab",
    "BIRDS OF THE WORLD URL": "Birds_of_the_World_URL",
    "AVIBASE ID": "AvibaseID",
    "ORIGINAL DESCRIPTION URL": "Original_description_URL",
}
REQUIRED_COLUMNS = ("Taxon_rank", "Scientific_name", "English_name_AviList",
                    "Species_code_Cornell_Lab")


def normalize_header(name) -> str:
    """Map an AviList header (any release) to the v2025 column name."""
    text = str(name or "").strip()
    if text in _HEADER_MAP:
        return _HEADER_MAP[text]
    if not text.isupper() and "_" in text:
        return text  # already v2025 style
    m = re.match(r"ENGLISH NAME: (CLEMENTS|BIRDLIFE) (v[\w.]+)$", text)
    if m:
        source = "Clements" if m.group(1) == "CLEMENTS" else "BirdLife"
        return f"English_name_{source}_{m.group(2)}"
    m = re.match(r"(v[\w.]+) CHANGES$", text)
    if m:
        return f"Changes_{m.group(1)}"
    words = text.split()
    if not words:
        return text
    return "_".join([words[0].capitalize(), *(w.lower() for w in words[1:])])


def download_xlsx(url: str) -> bytes:
    """Download the AviList XLSX file."""
    print(f"  Downloading {url}...")
    req = Request(url, headers={"User-Agent": USER_AGENT})
    try:
        with urlopen(req, timeout=120) as resp:
            data = resp.read()
    except (HTTPError, URLError, TimeoutError) as e:
        print(f"  ERROR: {e}")
        raise SystemExit(1)
    print(f"  Downloaded {len(data) / 1024 / 1024:.1f} MB")
    return data


def xlsx_to_csv(xlsx_bytes: bytes, csv_path: Path):
    """Convert XLSX to semicolon-delimited CSV."""
    try:
        import openpyxl
    except ImportError:
        print("ERROR: openpyxl is required. Install with: pip install openpyxl")
        raise SystemExit(1)

    print("  Converting XLSX to CSV...")
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
    # Releases ship several tabs (checklist, decision summaries, how-to, ...);
    # pick the one whose header carries the taxon rank column.
    ws = header = None
    for sheet in wb.worksheets:
        first = next(sheet.iter_rows(values_only=True, max_row=1), ())
        cols = [normalize_header(c) for c in first]
        if "Taxon_rank" in cols and "Scientific_name" in cols:
            ws, header = sheet, cols
            break
    if ws is None:
        print("ERROR: no checklist sheet with rank and scientific name columns found")
        raise SystemExit(1)
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        print(f"ERROR: AviList sheet '{ws.title}' lacks columns: {', '.join(missing)}")
        raise SystemExit(1)

    tmp_path = csv_path.with_suffix(".tmp")
    with open(tmp_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        writer.writerow(header)
        row_count = 0
        for row in ws.iter_rows(values_only=True, min_row=2):
            if not any(cell not in (None, "") for cell in row):
                continue
            writer.writerow([cell if cell is not None else "" for cell in row])
            row_count += 1
    tmp_path.replace(csv_path)

    wb.close()
    print(f"  Wrote {row_count} rows from sheet '{ws.title}' to {csv_path.name}")


def main():
    cfg = load_config()
    avilist_cfg = cfg.get("avilist", {})
    url = avilist_cfg.get("url", "")
    csv_name = avilist_cfg.get("csv_file", "")

    if not url or not csv_name:
        print("ERROR: avilist.url and avilist.csv_file must be set in config.yml")
        raise SystemExit(1)

    csv_path = RAW_DIR / csv_name

    parser = argparse.ArgumentParser(description="Download and convert AviList checklist")
    parser.add_argument("--force", action="store_true",
                        help="Re-download even if CSV already exists")
    args = parser.parse_args()

    if csv_path.exists() and not args.force:
        print(f"AviList CSV already exists: {csv_path.name}")
        print("  Use --force to re-download.")
        return

    RAW_DIR.mkdir(parents=True, exist_ok=True)

    xlsx_data = download_xlsx(url)

    # Also save the xlsx for reference
    xlsx_name = url.rsplit("/", 1)[-1]
    xlsx_path = RAW_DIR / xlsx_name
    xlsx_path.write_bytes(xlsx_data)
    print(f"  Saved XLSX as {xlsx_path.name}")

    xlsx_to_csv(xlsx_data, csv_path)
    print("Done!")


if __name__ == "__main__":
    main()
